# -*- coding: utf-8 -*-
"""共享 docx 生成工具：中文论文翻译稿统一格式。
一级标题宋体三号16pt、二级宋体四号14pt、三级宋体小四12pt、正文宋体小四12pt；
行距1.5倍；标题不缩进；正文首行缩进两格；中文宋体、西文Times New Roman。
"""
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def set_run(run, size, bold):
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts')
        rpr.append(rf)
    rf.set(qn('w:eastAsia'), '宋体')


def make_doc():
    doc = Document()
    for sec in doc.sections:
        sec.left_margin = Cm(2.8)
        sec.right_margin = Cm(2.8)
        sec.top_margin = Cm(2.54)
        sec.bottom_margin = Cm(2.54)
    sizes = {'Title': 22, 'Heading 1': 16, 'Heading 2': 14, 'Heading 3': 12}
    for nm, sz in sizes.items():
        st = doc.styles[nm]
        st.font.name = "Times New Roman"
        st.font.size = Pt(sz)
        st.font.bold = True
        rpr = st.element.get_or_add_rPr()
        # 中文字体
        rf = rpr.find(qn('w:rFonts'))
        if rf is None:
            rf = OxmlElement('w:rFonts')
            rpr.append(rf)
        rf.set(qn('w:eastAsia'), '宋体')
        # 去颜色、去边框
        color = rpr.find(qn('w:color'))
        if color is not None:
            rpr.remove(color)
        pPr = st.element.find(qn('w:pPr'))
        if pPr is not None:
            pbdr = pPr.find(qn('w:pBdr'))
            if pbdr is not None:
                pPr.remove(pbdr)
            ind = pPr.find(qn('w:ind'))
            if ind is not None:
                pPr.remove(ind)
        st.paragraph_format.line_spacing = 1.5
        st.paragraph_format.first_line_indent = Pt(0)
    return doc


def add_title(doc, text):
    p = doc.add_paragraph(style='Title')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p.add_run(text), 22, True)


def add_h1(doc, text):
    p = doc.add_paragraph(style='Heading 1')
    set_run(p.add_run(text), 16, True)


def add_h2(doc, text):
    p = doc.add_paragraph(style='Heading 2')
    set_run(p.add_run(text), 14, True)


def add_h3(doc, text):
    p = doc.add_paragraph(style='Heading 3')
    set_run(p.add_run(text), 12, True)


def add_body(doc, text):
    p = doc.add_paragraph()
    set_run(p.add_run(text), 12, False)
    pPr = p._p.get_or_add_pPr()
    ind = OxmlElement('w:ind')
    ind.set(qn('w:firstLineChars'), '200')
    ind.set(qn('w:firstLine'), '480')
    pPr.append(ind)
    p.paragraph_format.line_spacing = 1.5


def save(doc, path):
    doc.save(path)